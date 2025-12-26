#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
数据摄取与解析主管Agent

负责：
1. 接受用户上传的文件
2. 自动识别文件类型（CSV、PDF）
3. 解析文件内容
4. 打包解析结果，准备发送给后续处理模块
"""

from crewai import Agent, Task, Crew
from crewai.tools import tool
from agents.base import get_llm
import os
import logging
import pandas as pd
from typing import Dict, Any, Optional
from pathlib import Path
import fitz  # PyMuPDF
from agents.pdf_metadata_agent import extract_pdf_metadata
from utils.vector_db import extract_text_from_pdf
import docx
import openpyxl
import yaml
import json


@tool("文件类型识别工具")
def identify_file_type_tool(file_path: str) -> str:
    """
    识别文件类型。根据文件扩展名识别文件类型。
    输入参数：file_path (字符串，文件的完整路径)
    返回：文件类型字符串，如 'csv'、'pdf'、'txt'、'md'、'docx'、'doc'、'json'、'yaml'、'xls'、'xlsx' 或 'unknown'
    """
    try:
        file_ext = Path(file_path).suffix.lower()
        
        ext_map = {
            '.csv': 'csv',
            '.pdf': 'pdf',
            '.txt': 'txt',
            '.md': 'md',
            '.markdown': 'md',
            '.mdx': 'md',
            '.docx': 'docx',
            '.doc': 'doc',
            '.json': 'json',
            '.yaml': 'yaml',
            '.yml': 'yaml',
            '.xls': 'xls',
            '.xlsx': 'xlsx'
        }
        
        return ext_map.get(file_ext, 'unknown')
    except Exception as e:
        logging.error(f"识别文件类型失败: {str(e)}")
        return 'unknown'


@tool("通用文本文件解析工具")
def parse_text_file_tool(file_path: str) -> str:
    """
    解析纯文本文件（txt, md, json, yaml）。
    输入参数：file_path (字符串，文件的完整路径)
    返回：JSON格式字符串，包含文件内容和元数据
    """
    import json
    import os
    from pathlib import Path
    
    result = {
        'file_type': Path(file_path).suffix.lower().replace('.', ''),
        'content': '',
        'file_path': file_path,
        'file_name': os.path.basename(file_path),
        'metadata': {}
    }
    
    try:
        # 尝试多种编码格式读取
        encodings = ['utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'latin1']
        content = None
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        
        if content is None:
            raise Exception("无法识别文件编码")
            
        result['content'] = content
        
        # 如果是 JSON 或 YAML，尝试解析以获取更丰富的元数据
        if result['file_type'] == 'json':
            try:
                parsed = json.loads(content)
                if isinstance(parsed, dict):
                    result['metadata'] = parsed
            except:
                pass
        elif result['file_type'] in ['yaml', 'yml']:
            try:
                parsed = yaml.safe_load(content)
                if isinstance(parsed, dict):
                    result['metadata'] = parsed
            except:
                pass
                
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        logging.error(f"解析文本文件失败: {str(e)}")
        return json.dumps({'error': str(e), 'success': False})


@tool("Office文档解析工具")
def parse_office_doc_tool(file_path: str) -> str:
    """
    解析Office文档（docx, xlsx）。
    输入参数：file_path (字符串，文件的完整路径)
    返回：JSON格式字符串，包含文件内容和元数据
    """
    import json
    import os
    from pathlib import Path
    
    file_ext = Path(file_path).suffix.lower()
    result = {
        'file_type': file_ext.replace('.', ''),
        'content': '',
        'file_path': file_path,
        'file_name': os.path.basename(file_path),
        'metadata': {}
    }
    
    try:
        if file_ext == '.docx':
            doc = docx.Document(file_path)
            result['content'] = '\n'.join([para.text for para in doc.paragraphs])
        elif file_ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path, data_only=True)
            content = []
            for sheet in wb.worksheets:
                content.append(f"Sheet: {sheet.title}")
                for row in sheet.rows:
                    content.append('\t'.join([str(cell.value or '') for cell in row]))
            result['content'] = '\n'.join(content)
        else:
            # 对于 .doc 和 .xls，目前简单报错或提示需转换为新格式
            # 在 Linux 下可以使用 libreoffice 转换，但在 Windows 下较复杂
            raise Exception(f"暂不支持直接解析 {file_ext} 格式，请先转换为 .docx 或 .xlsx")
            
        return json.dumps(result, ensure_ascii=False, indent=2)
    except Exception as e:
        logging.error(f"解析Office文件失败: {str(e)}")
        return json.dumps({'error': str(e), 'success': False})


@tool("CSV文件解析工具")
def parse_csv_tool(csv_path: str) -> str:
    """
    解析CSV文件，通过首行（表头）确定每一列的内容。
    输入参数：csv_path (字符串，CSV文件的完整路径)
    返回：JSON格式字符串，包含列信息、数据行、总行数等
    """
    import json
    
    result = {
        'file_type': 'csv',
        'columns': [],
        'rows': [],
        'total_rows': 0,
        'encoding': None,
        'file_path': csv_path,
        'file_name': os.path.basename(csv_path)
    }
    
    try:
        # 尝试多种编码格式读取CSV文件
        encodings = ['utf-8-sig', 'utf-8', 'gbk', 'gb2312', 'gb18030', 'latin1', 'cp936']
        df = None
        used_encoding = None
        
        for encoding in encodings:
            try:
                df = pd.read_csv(csv_path, encoding=encoding)
                used_encoding = encoding
                logging.info(f"成功使用 {encoding} 编码读取CSV文件")
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
            except Exception as e:
                if 'Unicode' in str(type(e).__name__):
                    continue
                if encoding == encodings[0]:
                    raise
        
        if df is None:
            raise Exception("无法使用常见编码格式读取CSV文件")
        
        result['encoding'] = used_encoding
        result['total_rows'] = len(df)
        
        # 分析每一列
        for idx, col_name in enumerate(df.columns):
            # 获取该列的非空示例值（最多3个）
            sample_values = df[col_name].dropna().head(3).tolist()
            sample_values = [str(v)[:50] for v in sample_values]  # 限制示例值长度
            
            column_info = {
                'name': str(col_name),
                'index': idx,
                'sample': sample_values,
                'data_type': str(df[col_name].dtype),
                'non_null_count': int(df[col_name].notna().sum()),
                'null_count': int(df[col_name].isna().sum())
            }
            result['columns'].append(column_info)
        
        # 保存前10行数据作为示例（避免数据过大）
        max_sample_rows = 10
        sample_df = df.head(max_sample_rows)
        # 将NaN值替换为None（JSON兼容）
        sample_df = sample_df.where(pd.notna(sample_df), None)
        rows_dict = sample_df.to_dict('records')
        result['rows'] = rows_dict
        
        logging.info(f"CSV解析成功: {len(result['columns'])} 列, {result['total_rows']} 行")
        
    except Exception as e:
        logging.error(f"解析CSV文件失败: {str(e)}")
        result['error'] = str(e)
    
    # 处理NaN值，确保JSON兼容
    def clean_for_json(obj):
        """递归清理对象中的NaN值"""
        if isinstance(obj, dict):
            return {k: clean_for_json(v) for k, v in obj.items()}
        elif isinstance(obj, list):
            return [clean_for_json(item) for item in obj]
        elif isinstance(obj, float) and (pd.isna(obj) or pd.isnull(obj)):
            return None  # 将NaN转换为None（JSON中的null）
        elif pd.isna(obj):
            return None
        else:
            return obj
    
    cleaned_result = clean_for_json(result)
    return json.dumps(cleaned_result, ensure_ascii=False, indent=2)


@tool("PDF文件解析工具")
def parse_pdf_tool(pdf_path: str) -> str:
    """
    解析PDF文件，提取原始文本和元数据（标题、作者、期刊等）。
    输入参数：pdf_path (字符串，PDF文件的完整路径)
    返回：JSON格式字符串，包含原始文本、元数据、页数、文件大小等
    """
    import json
    
    result = {
        'file_type': 'pdf',
        'raw_text': '',
        'metadata': {},
        'page_count': 0,
        'file_size': 0,
        'file_path': pdf_path,
        'file_name': os.path.basename(pdf_path)
    }
    
    try:
        # 获取文件大小
        result['file_size'] = os.path.getsize(pdf_path)
        
        # 提取PDF文本
        try:
            raw_text = extract_text_from_pdf(pdf_path)
            result['raw_text'] = raw_text
            logging.info(f"PDF文本提取成功: {len(raw_text)} 字符")
        except Exception as e:
            logging.warning(f"PDF文本提取失败: {str(e)}")
            result['raw_text'] = ''
        
        # 获取页数
        try:
            pdf_doc = fitz.open(pdf_path)
            result['page_count'] = len(pdf_doc)
            pdf_doc.close()
        except Exception as e:
            logging.warning(f"获取PDF页数失败: {str(e)}")
            result['page_count'] = 0
        
        # 提取元数据（标题、作者、摘要等）
        try:
            metadata = extract_pdf_metadata(pdf_path)
            if metadata:
                result['metadata'] = {
                    'title': metadata.get('title', ''),
                    'authors': metadata.get('authors', ''),
                    'abstract': metadata.get('abstract', ''),
                    'keywords': metadata.get('keywords', ''),
                    'journal': metadata.get('journal', ''),
                    'year': metadata.get('year', ''),
                    'doi': metadata.get('doi', '')
                }
                logging.info(f"PDF元数据提取成功: 标题={metadata.get('title', '')[:50]}...")
            else:
                # 如果元数据提取失败，至少尝试从文件名提取标题
                file_name = os.path.basename(pdf_path)
                title_from_filename = os.path.splitext(file_name)[0]
                result['metadata'] = {
                    'title': title_from_filename,
                    'authors': '',
                    'abstract': '',
                    'keywords': '',
                    'journal': '',
                    'year': '',
                    'doi': ''
                }
                logging.warning(f"PDF元数据提取失败，使用文件名作为标题: {title_from_filename}")
        except Exception as e:
            logging.warning(f"PDF元数据提取失败: {str(e)}")
            # 使用文件名作为标题
            file_name = os.path.basename(pdf_path)
            title_from_filename = os.path.splitext(file_name)[0]
            result['metadata'] = {
                'title': title_from_filename,
                'authors': '',
                'abstract': '',
                'keywords': '',
                'journal': '',
                'year': '',
                'doi': ''
            }
        
        logging.info(f"PDF解析成功: {result['page_count']} 页, {result['file_size']} 字节")
        
    except Exception as e:
        logging.error(f"解析PDF文件失败: {str(e)}")
        result['error'] = str(e)
    
    return json.dumps(result, ensure_ascii=False, indent=2)


def create_data_parser_agent():
    """
    创建数据摄取与解析主管Agent
    
    Returns:
        Agent实例
    """
    return Agent(
        role="数据摄取与解析主管",
        goal=(
            "负责接收用户上传的文件或本地笔记文件，自动识别文件类型（CSV、PDF、TXT、MD、DOCX、JSON、YAML、XLSX），"
            "并调用相应的工具解析文件内容。提取文件的原始文本、结构化数据和元数据。"
            "最后将解析结果打包成标准JSON格式，确保包含内容字段（content）或解析结果。"
        ),
        backstory=(
            "你是一位专业的数据摄取与解析主管，擅长处理各种格式的学术数据和个人笔记文件。"
            "你能够快速识别文件类型，并根据文件类型选择合适的解析方法。"
            "你能够处理纯文本、结构化表格和富文本格式，确保提取的信息完整且准确。"
            "你的工作是为后续的知识工程处理提供高质量的结构化数据。"
        ),
        tools=[
            identify_file_type_tool,
            parse_csv_tool,
            parse_pdf_tool,
            parse_text_file_tool,
            parse_office_doc_tool
        ],
        allow_delegation=False,
        verbose=True,
        llm=get_llm(),
        max_iter=5,
        max_execution_time=600
    )


def create_data_parsing_task(file_path: str) -> Task:
    """
    创建数据解析任务
    
    Args:
        file_path: 文件路径
        
    Returns:
        Task实例
    """
    file_name = os.path.basename(file_path)
    
    # 确保文件路径使用正确的格式（Windows路径需要转义反斜杠）
    normalized_path = file_path.replace('\\', '/') if os.sep == '\\' else file_path
    
    return Task(
        description=(
            f"## 任务目标\n"
            f"解析用户上传或指定的笔记文件：{file_name}\n\n"
            f"## 文件路径\n"
            f"原始路径: {file_path}\n"
            f"标准化路径: {normalized_path}\n\n"
            f"## 执行步骤\n"
            f"1. **识别文件类型**：使用工具识别文件格式。\n"
            f"2. **调用解析工具**：\n"
            f"   - 如果是 CSV，使用 CSV 解析工具。\n"
            f"   - 如果是 PDF，使用 PDF 解析工具。\n"
            f"   - 如果是 TXT、MD、JSON、YAML，使用通用文本文件解析工具。\n"
            f"   - 如果是 DOCX、XLSX，使用 Office 文档解析工具。\n"
            f"3. **格式化输出**：确保输出包含文件的完整内容，如果是笔记类文件，内容应存储在 `content` 字段中。\n\n"
            f"## 输出要求\n"
            f"输出格式必须是JSON，包含以下字段：\n"
            f"- file_type: 文件类型\n"
            f"- parse_result: 详细解析结果（包含 content 或 rows）\n"
            f"- success: 是否成功（true/false）\n"
            f"- file_name: 文件名\n"
            f"- file_path: 文件路径\n"
        ),
        agent=None,  # 将在创建Crew时指定
        expected_output="JSON格式的解析结果，包含文件类型、解析数据（内容或表格行）和元数据"
    )


def parse_file_with_agent(file_path: str) -> Dict[str, Any]:
    """
    使用Agent解析文件
    """
    import json
    import re
    
    try:
        # 创建Agent和Task
        agent = create_data_parser_agent()
        task = create_data_parsing_task(file_path)
        task.agent = agent
        
        # 创建Crew并执行
        crew = Crew(
            agents=[agent],
            tasks=[task],
            verbose=False
        )
        
        result = crew.kickoff()
        
        # 提取结果文本
        result_text = result.raw if hasattr(result, 'raw') else str(result)
        
        # 尝试健壮解析 JSON
        def robust_json_parse(text):
            # 1. 尝试匹配 Markdown 代码块中的 JSON
            json_match = re.search(r'```json\s*(\{.*?\})\s*```', text, re.DOTALL)
            if json_match:
                try: return json.loads(json_match.group(1).strip())
                except: pass
            
            # 2. 尝试寻找最后一个 JSON 对象（Agent 喜欢在最后输出 JSON）
            # 或者寻找最外层的 JSON
            json_blocks = re.findall(r'(\{.*\})', text, re.DOTALL)
            if json_blocks:
                # 尝试从后往前解析，因为 Agent 往往在总结后给出一个 JSON
                for block in reversed(json_blocks):
                    try:
                        res = json.loads(block.strip())
                        if isinstance(res, dict) and ('parse_result' in res or 'content' in res or 'rows' in res):
                            return res
                    except: continue
            
            # 3. 实在没有 JSON，但包含文字，将其作为 content 返回（兜底）
            if len(text.strip()) > 50:
                return {
                    'success': True,
                    'file_type': Path(file_path).suffix.replace('.', ''),
                    'parse_result': {
                        'content': text.strip(),
                        'metadata': {'title': os.path.splitext(os.path.basename(file_path))[0]}
                    }
                }
            return None

        parse_result = robust_json_parse(result_text)
        
        if parse_result:
            # 规范化结果格式
            if 'parse_result' in parse_result:
                if isinstance(parse_result['parse_result'], str):
                    try: parse_result['parse_result'] = json.loads(parse_result['parse_result'])
                    except: pass
            return parse_result
        
        # 解析失败，回退到直接工具调用
        logging.warning(f"Agent 结果解析失败，回退到直接工具调用: {file_path}")
        return parse_file_directly(file_path)
        
    except Exception as e:
        logging.error(f"使用Agent解析文件失败: {str(e)}", exc_info=True)
        return parse_file_directly(file_path)


def parse_file_directly(file_path: str) -> Dict[str, Any]:
    """
    直接调用工具解析文件（不使用Agent，作为备选方案）
    
    Args:
        file_path: 文件路径
        
    Returns:
        解析结果字典
    """
    import json
    
    try:
        # 识别文件类型 - Tool对象需要使用.run()方法调用
        file_type_result = identify_file_type_tool.run(file_path)
        # 如果返回的是字符串，直接使用；如果是其他格式，提取字符串
        if isinstance(file_type_result, str):
            file_type = file_type_result
        else:
            file_type = str(file_type_result)
        
        logging.info(f"直接解析模式：识别到文件类型 {file_type} -> {file_path}")
        
        result_str = None
        if file_type in ['txt', 'md', 'json', 'yaml']:
            result_str = parse_text_file_tool.run(file_path)
        elif file_type in ['docx', 'xlsx']:
            result_str = parse_office_doc_tool.run(file_path)
        elif file_type == 'csv':
            result_str = parse_csv_tool.run(file_path)
        elif file_type == 'pdf':
            result_str = parse_pdf_tool.run(file_path)
        elif file_type in ['doc', 'xls']:
            # 兼容旧格式提示
            return {
                'file_type': file_type,
                'success': False,
                'error': f"暂不支持直接解析 {file_type} 格式，请先转换为 .docx 或 .xlsx",
                'file_path': file_path,
                'file_name': os.path.basename(file_path)
            }
        else:
            return {
                'file_type': 'unknown',
                'success': False,
                'error': f'不支持的文件类型: {Path(file_path).suffix}',
                'file_path': file_path,
                'file_name': os.path.basename(file_path)
            }
        
        if result_str:
            if isinstance(result_str, str):
                parse_result = json.loads(result_str)
            else:
                parse_result = result_str
                
            return {
                'file_type': file_type,
                'parse_result': parse_result,
                'success': True,
                'file_path': file_path,
                'file_name': os.path.basename(file_path)
            }
        else:
            raise Exception("解析工具未返回结果")
        
    except Exception as e:
        logging.error(f"直接解析文件失败: {str(e)}", exc_info=True)
        return {
            'file_type': 'unknown',
            'success': False,
            'error': str(e),
            'file_path': file_path,
            'file_name': os.path.basename(file_path)
        }


def prepare_for_delivery(parse_result: Dict[str, Any]) -> Dict[str, Any]:
    """
    打包解析结果，准备发送给后续处理模块
    
    Args:
        parse_result: 解析结果字典
        
    Returns:
        打包后的数据字典
    """
    from datetime import datetime
    
    # 提取实际的解析数据
    if 'parse_result' in parse_result:
        data = parse_result['parse_result']
    else:
        data = parse_result
    
    packaged_data = {
        'source_type': data.get('file_type', parse_result.get('file_type', 'unknown')),
        'data': data,
        'metadata': {
            'file_name': data.get('file_name', parse_result.get('file_name', '')),
            'file_path': data.get('file_path', parse_result.get('file_path', '')),
            'processed_at': datetime.now().isoformat()
        },
        'timestamp': datetime.now().isoformat(),
        'success': parse_result.get('success', True)
    }
    
    return packaged_data

